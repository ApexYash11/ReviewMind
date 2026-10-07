"""Unit tests for uploaded-file parsing (txt/csv/pdf/docx/xlsx)."""

import io

from app.main import read_uploaded_file


class FakeUpload:
    """Minimal stand-in for a Streamlit UploadedFile."""

    def __init__(self, name: str, data: bytes) -> None:
        self.name = name
        self._data = data

    def read(self) -> bytes:
        return self._data


def _make_pdf_bytes(text_lines: list[str]) -> bytes:
    """Build a minimal one-page PDF with valid xref offsets."""
    content = "BT /F1 12 Tf 72 720 Td 14 TL\n"
    for line in text_lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content += f"({escaped}) Tj T*\n"
    content += "ET"
    content_bytes = content.encode("latin-1")
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content_bytes)).encode("ascii") + b" >>\nstream\n"
        + content_bytes + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode("ascii") + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(bodies) + 1}\n0000000000 65535 f \n".encode("ascii")
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(bodies) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF".encode("ascii")
    )
    return out


def _make_docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _make_xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _make_blank_pdf_bytes() -> bytes:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# --- txt ---------------------------------------------------------------


def test_txt_parsed():
    result = read_uploaded_file(FakeUpload("r.txt", b"Great camera.\n\nPoor battery."))
    assert result is not None and "Great camera" in result


def test_txt_empty_rejected():
    assert read_uploaded_file(FakeUpload("r.txt", b"   \n  ")) is None


# --- csv ---------------------------------------------------------------


def test_csv_with_review_column():
    data = b"id,review,rating\n1,Great camera,5\n2,Battery is bad,1\n"
    result = read_uploaded_file(FakeUpload("r.csv", data))
    assert result is not None and "Great camera" in result and "Battery is bad" in result


def test_csv_without_review_column_joins_fields():
    data = b"alpha,beta\nGreat camera,5\n"
    result = read_uploaded_file(FakeUpload("r.csv", data))
    assert result is not None and "Great camera 5" in result


def test_csv_empty_rejected():
    assert read_uploaded_file(FakeUpload("r.csv", b"")) is None


def test_csv_header_only_rejected():
    assert read_uploaded_file(FakeUpload("r.csv", b"review,rating\n")) is None


# --- pdf ---------------------------------------------------------------


def test_pdf_text_extracted():
    pdf = _make_pdf_bytes(["Great camera quality.", "Battery drains fast."])
    result = read_uploaded_file(FakeUpload("r.pdf", pdf))
    assert result is not None and "Great camera" in result and "Battery" in result


def test_pdf_without_text_rejected():
    assert read_uploaded_file(FakeUpload("r.pdf", _make_blank_pdf_bytes())) is None


def test_pdf_corrupt_rejected():
    assert read_uploaded_file(FakeUpload("r.pdf", b"%PDF-1.4 not really")) is None


# --- docx --------------------------------------------------------------


def test_docx_paragraphs_joined():
    docx = _make_docx_bytes(["Great camera quality.", "Battery drains fast."])
    result = read_uploaded_file(FakeUpload("r.docx", docx))
    assert result is not None and "Great camera" in result and "Battery" in result


def test_docx_empty_rejected():
    assert read_uploaded_file(FakeUpload("r.docx", _make_docx_bytes([]))) is None


def test_docx_corrupt_rejected():
    assert read_uploaded_file(FakeUpload("r.docx", b"PK not a zip")) is None


# --- xlsx --------------------------------------------------------------


def test_xlsx_with_review_column():
    xlsx = _make_xlsx_bytes([["id", "review"], [1, "Great camera"], [2, "Battery is bad"]])
    result = read_uploaded_file(FakeUpload("r.xlsx", xlsx))
    assert result is not None and "Great camera" in result and "Battery is bad" in result


def test_xlsx_without_review_column_joins_fields():
    xlsx = _make_xlsx_bytes([["a", "b"], ["Great camera", 5]])
    result = read_uploaded_file(FakeUpload("r.xlsx", xlsx))
    assert result is not None and "Great camera" in result


def test_xlsx_corrupt_rejected():
    assert read_uploaded_file(FakeUpload("r.xlsx", b"not a workbook")) is None


# --- unsupported -------------------------------------------------------


def test_unsupported_type_rejected():
    assert read_uploaded_file(FakeUpload("r.json", b'{"a": 1}')) is None
